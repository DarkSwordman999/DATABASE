"""Проверка отчётов reports/docx по требованиям методичек ЛР1-ЛР8.
Запуск из корня проекта: python reports/check_reports.py
Для каждого отчёта проверяется наличие обязательных разделов и сведений, отсутствие пустых
листингов и соответствие номера варианта на титульном листе."""
import os
import sys

from docx import Document

DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docx")

COMMON = ["Цель работы", "Используемое программное обеспечение", "Выводы"]
REQUIRED = {
    1: ["pg_hba.conf", "брандмауэр", "s0.bat", "s1.bat", "s.bat", "s1.bat DATA\\create_DB",
        "psql.exe", "libpq.dll", "192.168.0.102", "Задание 1", "Задание 2", "Windows 11",
        "PostgreSQL 18.6"],
    2: ["Windows 11", "PostgreSQL 18.6", "add_data", "CURRENT_TIME", "\\timing on",
        "idx_names", "btree", "hash", "PRIMARY KEY", "EXPLAIN ANALYZE", "Протокол измерений"],
    3: ["CREATE TYPE", "DROP TYPE", "CREATE DOMAIN", "CREATE TEMPORARY TABLE",
        "CREATE OPERATOR", "Разработанные функции и операторы"],
    4: ["Windows 11", "PostgreSQL 18.6", "dump-1.bat", "dump-2.bat", "dump-3.bat", "task0-01",
        "task0-02", "task0-03", "base_save", "base_rar.rar", "base_rarM.part1.rar",
        "FC: no differences encountered", "run_all.bat", "app-pgdump", "app-psql", "rar"],
    5: ["Windows 11", "PostgreSQL 18.6", "Visual Studio 2022", "c1.bat", "dll.bat", "show1.bat",
        "PG_FUNCTION_INFO_V1", "STRICT", "Dump of file"],
    6: ["Windows 11", "PostgreSQL 18.6", "SQL Server 2019", "s_TCP.bat", "pg_dump_str.bat",
        "pg_dump_data.bat", "sel2.cpp", "BULK INSERT", "copy_to_MS_SQL.bat", "disp.bat",
        "Задание 1", "Задание 2", "Последовательность действий"],
    7: ["create_objects.sql", "calculate1.sql", "calculate2.sql", "CREATE OR ALTER VIEW",
        "CREATE OR ALTER FUNCTION", "Параметры", "59206.10", "общая сумма премии",
        "общая сумма затрат на хранение"],
    8: ["Доступ к СУБД MS SQL Server", "Структура запроса", "Текст программы",
        "Порядок использования параметров", "Результаты"],
}
VARIANT = {
    (1, 20): ["tasks\\v20_task1.sql", "категория товара"],
    (1, 22): ["tasks\\v22_task1.sql", "поставщик"],
    (2, 20): ["ТОВАР"], (2, 22): ["МАГАЗИН"],
    (3, 20): ["vector3"], (3, 22): ["rational"],
    (5, 20): ["f1_floor", "f2_alltrim"], (5, 22): ["f1_cosd", "f2_pos"],
    (6, 20): ["v20_task1.sql", "v20_task2.sql"], (6, 22): ["v22_task1.sql", "v22_task2.sql"],
    (4, 20): ["v20_task1.sql", "run_all.bat 20"], (4, 22): ["v22_task1.sql", "run_all.bat 22"],
    (7, 20): ["Андрей", "плащ"], (7, 22): ["Ольга", "шкаф"],
    (8, 20): ["Lab08.cs", "SqlConnection", "C#"], (8, 22): ["lab08.py", "pyodbc", "Python"],
}


def text_of(doc):
    parts = [p.text for p in doc.paragraphs]
    empty_listings = 0
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
                if len(t.columns) == 1 and not cell.text.strip():
                    empty_listings += 1
    return "\n".join(parts), empty_listings


def main():
    bad = 0
    for lab in range(1, 9):
        for v in (20, 22):
            name = f"ЛР{lab}_вариант_{v}.docx"
            path = os.path.join(DIR, name)
            if not os.path.exists(path):
                print(f"НЕТ ФАЙЛА  {name}")
                bad += 1
                continue
            doc = Document(path)
            text, empty = text_of(doc)
            missing = [k for k in COMMON + REQUIRED[lab] + VARIANT.get((lab, v), [])
                       if k not in text]
            if f"Вариант {v}" not in text:
                missing.append(f"титул: Вариант {v}")
            other = 22 if v == 20 else 20
            if f"Вариант {other}" in text.split("Лабораторная работа")[0]:
                missing.append("титул: чужой вариант")
            problems = missing + ([f"пустых листингов: {empty}"] if empty else [])
            listings = text.count("Листинг ")
            status = "OK  " if not problems else "FAIL"
            bad += bool(problems)
            print(f"{status} {name:22} листингов {listings:3}  "
                  + ("; ".join(problems) if problems else ""))
    print("Итог:", "все отчёты соответствуют требованиям" if not bad else f"проблем: {bad}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
