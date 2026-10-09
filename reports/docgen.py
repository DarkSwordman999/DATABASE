"""Оформление отчётов по ЛР (python-docx): титульный лист ТулГУ, стили, листинги, таблицы."""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "reports", "out")

DISCIPLINE = "Программирование и администрирование баз данных"
TEACHER = "Скобельцын С.А."
TEACHER_POS = "доцент, канд. техн. наук"
CITY_YEAR = "Тула 2026"

MONO = "Courier New"
SERIF = "Times New Roman"


def read(path):
    """Текст файла проекта (путь от корня репозитория)."""
    with open(os.path.join(ROOT, path), encoding="utf-8-sig") as f:
        return f.read().rstrip("\n")


def out(name):
    """Сохранённый вывод прогона reports/out/<name>.txt без строки вызова."""
    with open(os.path.join(OUT, name + ".txt"), encoding="utf-8") as f:
        text = f.read().rstrip("\n")
    if text.startswith(("> ./h ", "> ./help ")):  # строка вызова - она уже есть в подписи
        text = text.split("\n", 1)[1] if "\n" in text else ""
    return text.strip("\n")


def out_cmd(name):
    """Команда ./h или ./help, которой получен вывод reports/out/<name>.txt (или None)."""
    with open(os.path.join(OUT, name + ".txt"), encoding="utf-8") as f:
        first = f.readline().rstrip("\n")
    return first[2:] if first.startswith(("> ./h ", "> ./help ")) else None


# сценарии запуска: h - Git Bash, help - PowerShell (help.cmd -> helper/help.ps1)
SCRIPT = {"h": "h", "help": "helper/help.ps1"}


def _help_lines(runner):
    """Строки справки: echo "..." в h (None - не echo), блоки $Help.<блок> = @' ... '@ в help.ps1
    (None - граница блока)."""
    src = read(SCRIPT[runner]).splitlines()
    if runner == "h":
        for line in src:
            m = re.match(r'\s*echo "(.*)"$', line)
            yield m.group(1).replace("\\\\", "\\") if m else None
        return
    on = False                  # внутри блока $Help.<блок> = @' ... '@
    for line in src:
        if on and line.startswith("'@"):
            on = False
            yield None
        elif on:
            yield line
        elif re.match(r"\$Help\.\w+ = @'$", line.rstrip()):
            on = True


def h_help(lab, runner="h"):
    """Раздел справки ./h (./help) short <lab> - строки блока ЛР lab с примерами запуска."""
    lines, on = [], False
    for text in _help_lines(runner):
        if text is None:
            if on:
                break
            continue
        if text.startswith("=================="):
            if on:
                break
            on = f"ЛР{lab}:" in text
        if on:
            lines.append(text)
    return "\n".join(l for l in lines if l.strip())


def h_branch(lab, runner="h"):
    """Ветвь case (switch) сценария h (help.ps1), обрабатывающая команды ЛР lab."""
    lines = read(SCRIPT[runner]).splitlines()
    start = next(i for i, l in enumerate(lines) if re.match(rf"\s*# -+ ЛР{lab} -+", l))
    end = start + 1
    while end < len(lines) and not re.match(r"\s*# -+ ЛР\d -+|\s*'?reports'?\s*[){]",
                                            lines[end]):
        end += 1
    while lines[end - 1].strip() == "":
        end -= 1
    return "\n".join(lines[start:end])


def _set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


class Report:
    def __init__(self, lab, topic, variant, runner="help"):
        self.lab, self.topic, self.variant = lab, topic, variant
        self.runner = runner        # сценарий запуска: h (Git Bash) или help (PowerShell)
        self.fig = 0
        self.tab = 0
        self.lst = 0
        self.doc = Document()
        self._styles()
        self._title_page()

    # ------------------------------------------------------------------ стили
    def _styles(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)

        normal = self.doc.styles["Normal"]
        normal.font.name = SERIF
        normal.font.size = Pt(14)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        pf.line_spacing = 1.15

        for name, size in (("Heading 1", 14), ("Heading 2", 14)):
            st = self.doc.styles[name]
            st.font.name = SERIF
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = RGBColor(0, 0, 0)
            rpr = st.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(attr), SERIF)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rfonts.attrib.pop(qn(attr), None)     # шрифт темы перекрывает явный
            st.paragraph_format.space_before = Pt(12)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True

    def _center(self, text="", size=None, bold=False, space_after=0):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if text:
            r = p.add_run(text)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def _title_page(self):
        c = self._center
        c("МИНОБРНАУКИ РОССИИ")
        c()
        c("Федеральное государственное бюджетное образовательное учреждение")
        c("высшего образования")
        c("«Тульский государственный университет»")
        c()
        c("Институт прикладной математики и компьютерных наук")
        for _ in range(5):
            c()
        c("ОТЧЁТ", size=17, bold=True)
        c(f"по лабораторной работе №{self.lab}", size=15)
        c(f"по дисциплине «{DISCIPLINE}»")
        c()
        c(f"«{self.topic}»", bold=True)
        c()
        c(f"Вариант {self.variant}")
        for _ in range(5):
            c()
        for left, right in ((f"Выполнил: ст. гр. ________", "________________"),
                            (f"Принял: {TEACHER_POS}", TEACHER)):
            p = self.doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), alignment=2)  # RIGHT
            p.add_run(f"{left}\t{right}")
        for _ in range(8):
            c()
        c(CITY_YEAR)
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --------------------------------------------------------------- элементы
    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def p(self, text, bold_prefix=None):
        """Абзац с отступом первой строки; **жирный** выделяется двумя звёздочками."""
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.first_line_indent = Cm(1.25)
        if bold_prefix:
            par.add_run(bold_prefix).bold = True
        for i, part in enumerate(re.split(r"\*\*", text)):
            run = par.add_run(part)
            run.bold = i % 2 == 1
        return par

    def bullets(self, items):
        for item in items:
            par = self.doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.left_indent = Cm(1.25)
            par.paragraph_format.first_line_indent = Cm(-0.5)
            for i, part in enumerate(re.split(r"\*\*", "– " + item)):
                run = par.add_run(part)
                run.bold = i % 2 == 1

    def code(self, text, caption=None, size=9):
        """Моноширинный блок в рамке (листинг сценария или вывод программы)."""
        if caption:
            self.lst += 1
            cap = self.doc.add_paragraph()
            cap.paragraph_format.space_before = Pt(6)
            cap.paragraph_format.keep_with_next = True
            r = cap.add_run(f"Листинг {self.lst} – {caption}")
            r.italic = True
            r.font.size = Pt(12)
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        _set_cell_shading(cell, "F5F5F5")
        cell.paragraphs[0].text = ""
        lines = text.replace("\t", "    ").split("\n")
        first = True
        for line in lines:
            par = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            par.paragraph_format.line_spacing = 1.0
            run = par.add_run(line if line else " ")
            run.font.name = MONO
            run.font.size = Pt(size)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def listing(self, path, caption=None, size=9):
        self.code(read(path), caption or f"файл {path}", size)

    def output(self, name, caption, size=8.5):
        """Вывод прогона; в подписи - команда ./h (./help), которой он получен, и аналог в cmd."""
        cmd = out_cmd(name)
        if cmd:
            m = re.match(r"(.*?):\s*>(.*)$", caption)
            caption = (f"{m.group(1)}: {cmd} (аналог в cmd: >{m.group(2)})" if m
                       else f"{caption}: {cmd}")
        self.code(out(name), caption, size)

    def defense(self, num, commands, notes=()):
        """Раздел «Защита работы»: команды ./h (./help), справка и фрагмент сценария для ЛР."""
        cmd = f"./{self.runner}"
        self.h1(f"{num}. Защита работы: запуск через {cmd}")
        if self.runner == "help":
            self.p("Все действия работы выполняются из корня проекта в PowerShell (или cmd) "
                   "единым сценарием **help**: командный файл help.cmd вызывает "
                   "helper/help.ps1 с ключом -ExecutionPolicy Bypass, поэтому запуск не "
                   "зависит от политики выполнения сценариев PowerShell (листинги - в "
                   "приложении А). Сценарий задаёт параметры подключения к PostgreSQL "
                   "(PGHOST, PGPORT, PGUSER, PGDATABASE, клиентская кодировка UTF8), передаёт "
                   "параметры сценариям psql через стандартный ввод командами \\set "
                   "(кириллица в аргументах командной строки Windows не искажается), вызывает "
                   "командные файлы .bat и приводит служебные сообщения psql из CP1251 к "
                   "UTF-8 (функция Fix-Line). Git Bash для работы не требуется. Все "
                   "результаты, приведённые в отчёте, получены этими командами; в подписи "
                   "каждого листинга результата указана команда ./help.")
        else:
            self.p("Все действия работы выполняются из корня проекта в Git Bash единым "
                   "сценарием **h** (листинг полностью - в приложении А). Сценарий задаёт "
                   "параметры подключения к PostgreSQL (PGHOST, PGPORT, PGUSER, PGDATABASE, "
                   "клиентская кодировка UTF8), передаёт параметры сценариям psql через "
                   "стандартный ввод командами \\set (кириллица в аргументах командной строки "
                   "Windows не искажается), вызывает командные файлы .bat и приводит вывод к "
                   "UTF-8 фильтром helper/fixenc.pl. Все результаты, приведённые в отчёте, "
                   "получены этими командами; в подписи каждого листинга результата указана "
                   "команда ./h.")
        self.p("Порядок демонстрации работы на защите:")
        self.table(["№", "Команда", "Что выполняется"],
                   [(i, c, d) for i, (c, d) in enumerate(commands, 1)],
                   f"Команды {cmd} для защиты ЛР{self.lab}, вариант {self.variant}",
                   widths=[1, 7.5, 8], size=11)
        for n in notes:
            self.p(n)
        self.code(h_help(self.lab, self.runner),
                  f"справка по командам ЛР{self.lab}: {cmd} short {self.lab}", size=8)
        self.code(h_branch(self.lab, self.runner),
                  f"обработка команд ЛР{self.lab} в сценарии {SCRIPT[self.runner]} (фрагмент)",
                  size=8)

    def h_appendix(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if self.runner == "help":
            self.h1("Приложение А. Листинг сценария help")
            self.p("Сценарий help - единая точка запуска всех лабораторных работ из PowerShell "
                   "или cmd: ./help <команда> [параметры]; ./help без параметров выводит "
                   "список команд. Командный файл help.cmd запускает helper/help.ps1.")
            self.listing("help.cmd", "командный файл help.cmd", size=7.5)
            self.listing("helper/help.ps1", "сценарий helper/help.ps1", size=7.5)
            return
        self.h1("Приложение А. Листинг сценария h")
        self.p("Сценарий h - единая точка запуска всех лабораторных работ из Git Bash: "
               "./h <команда> [параметры]; ./h без параметров выводит список команд.")
        self.listing("h", "сценарий h", size=7.5)

    def table(self, header, rows, caption=None, widths=None, size=12):
        if caption:
            self.tab += 1
            cap = self.doc.add_paragraph()
            cap.paragraph_format.space_before = Pt(6)
            cap.paragraph_format.keep_with_next = True
            r = cap.add_run(f"Таблица {self.tab} – {caption}")
            r.font.size = Pt(12)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            _set_cell_shading(cell, "E7E6E6")
            run = cell.paragraphs[0].add_run(str(h))
            run.bold = True
            run.font.size = Pt(size)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                run = cells[i].paragraphs[0].add_run(str(v))
                run.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def software(self, items):
        self.h1("Используемое программное обеспечение")
        self.table(["Компонент", "Версия"], items, widths=[7, 9.5])

    def save(self, name):
        folder = os.path.join(ROOT, "reports", "docx")
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, name)
        self.doc.save(path)
        return path


OS = ("Операционная система", "Windows 11 Pro (сборка 10.0.22000)")
PG = ("Сервер PostgreSQL", "PostgreSQL 18.6 (x86_64-windows, msvc-19.44)")
PSQL = ("Клиент", "psql 18.6")
MSSQL = ("Сервер MS SQL Server", "Microsoft SQL Server 2019 (15.0.4153.1) Express LocalDB")
SQLCMD = ("Клиент MS SQL Server", "sqlcmd (ODBC Driver 17 for SQL Server)")
MSVC = ("Компилятор C/C++", "Microsoft C/C++ 19.33.31629 (Visual Studio 2022), Windows SDK 10.0.19041.0")
RAR = ("Архиватор", "RAR 5.80 (Rar.exe из WinRAR)")
